"""
Actualiza módulos (áreas generales) desde «ASIGNACIONES DE LK Y DT (56).xlsx».

Incluye: lockers, dotaciones, seca botas, asignaciones, personal, personal presupuestado.
NO toca: Historial de retiros (hoja RETIROS) ni planta Desposte.

Uso (desde la raíz del proyecto):
  python scripts/actualizar_desde_excel_lk_dt.py
  python scripts/actualizar_desde_excel_lk_dt.py --excel "ASIGNACIONES DE LK Y DT (56).xlsx"
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(SCRIPTS))

from dotenv import load_dotenv

load_dotenv(PROJECT_ROOT / ".env")

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Instala con: pip install openpyxl")
    sys.exit(1)

from app import create_app, db
from app.models import (
    BaseDotaciones,
    BaseLockers,
    HistorialRetiros,
    PersonalPresupuestado,
    RegistroAsignaciones,
    RegistroPersonal,
    SecaBotasDisponibles,
)
from import_datos import (
    IMPORTERS,
    _normalize_area_registro_asignaciones_csv,
)

# RETIROS queda fuera a propósito (Historial de retiros).
HOJAS_A_IMPORTAR = [
    ("LOCKERES", "base_lockers"),
    ("DOTACIONES", "base_dotaciones"),
    ("SECA BOTAS", "seca_botas_disponibles"),
    ("PERSONAL PRESUPUESTADO", "personal_presupuestado"),
    ("ASIGNACIONES", "registro_asignaciones"),
    ("PERSONAL", "registro_personal"),
]

DEFAULT_EXCEL = PROJECT_ROOT / "ASIGNACIONES DE LK Y DT (56).xlsx"


def _cell_to_str(cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(cell, bool):
        return str(cell)
    if isinstance(cell, int):
        return str(cell)
    if isinstance(cell, float):
        if cell.is_integer():
            return str(int(cell))
        return str(cell)
    return str(cell).strip()


def _sheet_to_rows(ws) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append([_cell_to_str(c) for c in row])
    if not rows:
        return rows
    last = 0
    for i, h in enumerate(rows[0]):
        if (h or "").strip():
            last = i
    width = last + 1
    return [r[:width] for r in rows]


def _norm_doc(val) -> str:
    return re.sub(r"\D+", "", str(val or "").strip())


def _hdr_index(hdr: list[str], *names: str) -> int | None:
    want = {n.strip().lower() for n in names}
    for i, h in enumerate(hdr):
        if (h or "").strip().lower() in want:
            return i
    return None


def _fila_con_datos(row: list[str]) -> bool:
    return any((c or "").strip() for c in row)


def _fila_asignacion_util(row: list[str], hdr: list[str]) -> bool:
    for name in ("operario", "identificacion", "codigo de lockets", "codigo de dotacion", "area"):
        i = _hdr_index(hdr, name)
        if i is not None and i < len(row) and (row[i] or "").strip():
            return True
    return False


def _docs_en_asignaciones(rows: list[list[str]]) -> set[str]:
    if not rows:
        return set()
    hdr = rows[0]
    i_doc = _hdr_index(hdr, "identificacion")
    if i_doc is None:
        return set()
    out: set[str] = set()
    for row in rows[1:]:
        if not _fila_asignacion_util(row, hdr):
            continue
        d = _norm_doc(row[i_doc] if i_doc < len(row) else "")
        if d:
            out.add(d)
    return out


def _contar_filas_utiles(rows: list[list[str]], hoja: str) -> int:
    if not rows or len(rows) < 2:
        return 0
    if hoja == "SECA BOTAS":
        i_cod = _hdr_index(rows[0], "codigo de seca botas", "codigo")
        n = 0
        for row in rows[1:]:
            if i_cod is not None and i_cod < len(row) and (row[i_cod] or "").strip():
                n += 1
        return n
    if hoja == "ASIGNACIONES":
        return sum(1 for row in rows[1:] if _fila_asignacion_util(row, rows[0]))
    if hoja == "PERSONAL":
        i_op = _hdr_index(rows[0], "operario")
        i_doc = _hdr_index(rows[0], "identificacion")
        n = 0
        for row in rows[1:]:
            op = row[i_op] if i_op is not None and i_op < len(row) else ""
            doc = row[i_doc] if i_doc is not None and i_doc < len(row) else ""
            if (op or "").strip() or (doc or "").strip():
                n += 1
        return n
    return sum(1 for row in rows[1:] if _fila_con_datos(row))


def _agregar_personal_pendiente(personal_rows: list[list[str]], docs_asig: set[str], app) -> int:
    """PERSONAL activo sin fila en ASIGNACIONES → Personal pendiente."""
    if not personal_rows or len(personal_rows) < 2:
        return 0
    hdr = personal_rows[0]
    i_id = _hdr_index(hdr, "id personal")
    i_op = _hdr_index(hdr, "operario")
    i_doc = _hdr_index(hdr, "identificacion")
    i_area = _hdr_index(hdr, "area")
    i_talla = _hdr_index(hdr, "talla")
    i_al = _hdr_index(hdr, "area de lockers")
    i_est = _hdr_index(hdr, "estado")

    added = 0
    with app.app_context():
        existentes = {
            _norm_doc(r.identificacion)
            for r in RegistroAsignaciones.query.filter(
                RegistroAsignaciones.es_planta_desposte.is_(False)
            ).all()
            if _norm_doc(r.identificacion)
        }
        for row in personal_rows[1:]:

            def g(i):
                return (row[i] if i is not None and i < len(row) else "").strip()

            operario = g(i_op)
            doc_raw = g(i_doc)
            doc = _norm_doc(doc_raw)
            if not operario and not doc:
                continue
            estado = g(i_est).upper()
            if estado and estado not in ("ACTIVO", "A", "1", "TRUE", "SI", "SÍ"):
                continue
            if not doc or doc in docs_asig or doc in existentes:
                continue
            area_raw = g(i_area)
            area = _normalize_area_registro_asignaciones_csv(area_raw) or area_raw
            if (area or "").strip().upper() == "DESPOSTE":
                continue
            db.session.add(
                RegistroAsignaciones(
                    id_asignaciones=(g(i_id) or "")[:50],
                    codigo_dotacion="",
                    fecha_asignacion=datetime.utcnow(),
                    fecha_entrega=None,
                    operario=operario[:120],
                    codigo_lockets="",
                    identificacion=(doc_raw or doc)[:40],
                    codigo_seca_botas="",
                    area=(area or "")[:100],
                    talla_operarios=g(i_talla)[:20],
                    talla_dotacion="",
                    area_lockers=g(i_al)[:100],
                    estado="Activo",
                    observaciones="",
                    es_planta_desposte=False,
                )
            )
            existentes.add(doc)
            added += 1
        if added:
            db.session.commit()
    print(f"personal_pendiente(desde PERSONAL): agregados {added}.")
    return added


def _validar(app, excel_counts: dict) -> bool:
    ok = True
    with app.app_context():
        pairs = [
            (
                "base_lockers (no DESPOSTE)",
                BaseLockers.query.filter(db.func.upper(db.func.trim(BaseLockers.area)) != "DESPOSTE").count(),
                excel_counts.get("base_lockers"),
            ),
            (
                "base_dotaciones (no DESPOSTE)",
                BaseDotaciones.query.filter(
                    db.func.upper(db.func.trim(BaseDotaciones.area_uso)) != "DESPOSTE"
                ).count(),
                excel_counts.get("base_dotaciones"),
            ),
            (
                "seca_botas_disponibles",
                SecaBotasDisponibles.query.count(),
                excel_counts.get("seca_botas_disponibles"),
            ),
            (
                "personal_presupuestado",
                PersonalPresupuestado.query.count(),
                excel_counts.get("personal_presupuestado"),
            ),
            (
                "registro_personal",
                RegistroPersonal.query.count(),
                excel_counts.get("registro_personal"),
            ),
            (
                "registro_asignaciones (no desposte)",
                RegistroAsignaciones.query.filter(RegistroAsignaciones.es_planta_desposte.is_(False)).count(),
                excel_counts.get("registro_asignaciones_total"),
            ),
            (
                "historial_retiros (intactos)",
                HistorialRetiros.query.filter(HistorialRetiros.es_planta_desposte.is_(False)).count(),
                excel_counts.get("historial_retiros_antes"),
            ),
            (
                "desposte asignaciones (intactos)",
                RegistroAsignaciones.query.filter(RegistroAsignaciones.es_planta_desposte.is_(True)).count(),
                excel_counts.get("desposte_asig_antes"),
            ),
            (
                "desposte lockers (intactos)",
                BaseLockers.query.filter(db.func.upper(db.func.trim(BaseLockers.area)) == "DESPOSTE").count(),
                excel_counts.get("desposte_lock_antes"),
            ),
            (
                "desposte dotaciones (intactos)",
                BaseDotaciones.query.filter(
                    db.func.upper(db.func.trim(BaseDotaciones.area_uso)) == "DESPOSTE"
                ).count(),
                excel_counts.get("desposte_dot_antes"),
            ),
        ]
        print("\n=== Validación ===")
        for label, got, expected in pairs:
            if expected is None:
                print(f"  [?] {label}: BD={got}")
                continue
            if got == expected:
                mark = "OK"
            elif label.startswith("registro_asignaciones") and got >= (
                excel_counts.get("registro_asignaciones") or 0
            ):
                mark = "OK+"
            else:
                mark = "DIFF"
                ok = False
            print(f"  [{mark}] {label}: BD={got} excel/antes={expected}")
    return ok


def main():
    parser = argparse.ArgumentParser(
        description="Actualizar BD desde Excel LK/DT (sin Historial de retiros ni Desposte)"
    )
    parser.add_argument("--excel", default=str(DEFAULT_EXCEL), help="Ruta al Excel")
    args = parser.parse_args()
    path = Path(args.excel)
    if not path.exists():
        print(f"No se encuentra el Excel: {path}")
        sys.exit(1)

    app = create_app()
    excel_counts: dict = {}

    with app.app_context():
        excel_counts["historial_retiros_antes"] = HistorialRetiros.query.filter(
            HistorialRetiros.es_planta_desposte.is_(False)
        ).count()
        excel_counts["desposte_asig_antes"] = RegistroAsignaciones.query.filter(
            RegistroAsignaciones.es_planta_desposte.is_(True)
        ).count()
        excel_counts["desposte_lock_antes"] = BaseLockers.query.filter(
            db.func.upper(db.func.trim(BaseLockers.area)) == "DESPOSTE"
        ).count()
        excel_counts["desposte_dot_antes"] = BaseDotaciones.query.filter(
            db.func.upper(db.func.trim(BaseDotaciones.area_uso)) == "DESPOSTE"
        ).count()

    print(f"Leyendo {path.name} ...")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    available = set(wb.sheetnames)
    if "RETIROS" in available:
        print("Omitido (pedido): hoja RETIROS / Historial de retiros.")

    sheets: dict[str, list[list[str]]] = {}
    for hoja, _tabla in HOJAS_A_IMPORTAR:
        if hoja in available:
            sheets[hoja] = _sheet_to_rows(wb[hoja])
    wb.close()

    personal_rows = sheets.get("PERSONAL") or []
    asig_rows = sheets.get("ASIGNACIONES") or []
    docs_asig = _docs_en_asignaciones(asig_rows)
    excel_counts["registro_asignaciones"] = _contar_filas_utiles(asig_rows, "ASIGNACIONES")

    for hoja, tabla in HOJAS_A_IMPORTAR:
        rows = sheets.get(hoja)
        if not rows or len(rows) < 2:
            print(f"[{tabla}] Hoja '{hoja}' vacia o ausente - omitida.")
            continue
        utiles = _contar_filas_utiles(rows, hoja)
        excel_counts[tabla] = utiles
        print(f"\n>>> {hoja} -> {tabla} ({utiles} filas utiles, replace)")
        fn = IMPORTERS.get(tabla)
        if not fn:
            print(f"  Sin importador para {tabla}")
            continue
        fn(rows, True, app)

    pend = _agregar_personal_pendiente(personal_rows, docs_asig, app)
    excel_counts["registro_asignaciones_total"] = excel_counts.get("registro_asignaciones", 0) + pend

    ok = _validar(app, excel_counts)
    print("\nListo." if ok else "\nListo con diferencias - revisa validacion.")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
